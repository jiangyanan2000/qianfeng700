import openpyxl
from docx import Document
import os
filelist = os.listdir(r"C:\Users\Admin\Desktop\村庄")
print(filelist)
wb = openpyxl.Workbook()
ws = wb.active
for file in filelist:
    doc = Document(r"C:\Users\Admin\Desktop\村庄"+"\\"+file)
    tables = doc.tables
    for i in range(len(tables)):
        tb = tables[i]
        tb_rows = tb.rows
        for i in range(1,len(tb_rows)):
            row_data = []
            row_cells = tb_rows[i].cells
            for cell in row_cells:
                if cell.text != " ":
                    row_data.append(cell.text)
                    row_data[0] = file.split(".")[0]
            print(row_data)
            ws.append(row_data[0:2])
wb.save("Huizong.xlsx")



